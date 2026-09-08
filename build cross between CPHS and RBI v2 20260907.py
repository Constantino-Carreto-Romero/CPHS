# -*- coding: utf-8 -*-
"""
Build the CPHS <-> RBI district crosswalk.

Maps CPHS district names onto 2011 Census (PC11) district codes, so the RBI
branch shares can be merged onto the CPHS household panel. Matching runs in
four layers: exact against SHRUG (2011 names), exact against the Local
Government Directory (current names), fuzzy candidates against LGD, and a
manual table for the residual.

--------------------------------------------------------------------
DEPENDENCIES
--------------------------------------------------------------------
Two packages are not part of the base Anaconda install. Install them from the
ANACONDA PROMPT (not from the Spyder console), then restart the Spyder kernel
so pandas picks them up:

    pip install python-calamine
    pip install rapidfuzz

  python-calamine  needed to read the LGD file. The lgdirectory.gov.in portal
                   exports .xlsx with malformed style definitions that the
                   default engine (openpyxl) rejects with a TypeError.
  rapidfuzz        used in Layer 3 to propose candidate matches by string
                   similarity.
--------------------------------------------------------------------
"""

from pathlib import Path
import pandas as pd
import pyarrow.dataset as ds

# ---------------------------------------------------------------
# Paths
# ---------------------------------------------------------------
BASE        = Path(r"C:\Users\HP\Documents\QMUL\Financial inclusion in India")
DATA        = BASE / "data"
TABLES      = BASE / "tables"
PARQUET_DIR = DATA / "CPHS" / "cphs_panel_parquet"
SHRID_KEYS  = DATA / "shrug-shrid-keys-csv"     # shrid_loc_names.csv
PC_KEYS     = DATA / "shrug-pc-keys-csv"        # shrid_pc11dist_key.csv

# TABLES.mkdir(parents=True, exist_ok=True)


##########################################################
# CPHS: Use parquet files to find all possible state-district combinations
# on the CPHS data
##########################################################

# Sanity check: how many parts did we find?
files = sorted(PARQUET_DIR.glob("part_*.parquet"))
print(f"Files found: {len(files)}")          # expect 144

# Read the three columns needed (month_date is required for the month count)
dataset = ds.dataset(PARQUET_DIR, format="parquet")
cphs = dataset.to_table(columns=["state", "district", "month_date"]).to_pandas()
print(f"Rows read: {len(cphs):,}")

# Unique month-state-district combinations, then count months per pair
cphs_pairs = (cphs[["state", "district", "month_date"]]
                .drop_duplicates()                       # one row per month-state-district
                .groupby(["state", "district"])
                .size()
                .reset_index(name="n_months")            # months in which the pair appears
                .sort_values(["state", "district"])
                .reset_index(drop=True))

print(f"Unique state-district pairs: {len(cphs_pairs)}")
print(cphs_pairs.head(20))

# Save pairs
out_path = TABLES / "CPHS state-district pairs.xlsx"
cphs_pairs.to_excel(out_path, index=False)
print(f"Saved {len(cphs_pairs)} pairs to: {out_path}")


### Sanity check

# Districts covered per month: is the panel's geographic coverage stable?
per_month = (cphs[["state", "district", "month_date"]]
               .drop_duplicates()
               .groupby("month_date")
               .size())
print(per_month.describe())
print(per_month.head(12))
print(per_month.tail(12))


#####################################################################
# RBI: check district divisions
#####################################################################

# Load the SHRUG keys once; both are reused further down.
names = pd.read_csv(SHRID_KEYS / "shrid_loc_names.csv", dtype=str)
codes = pd.read_csv(PC_KEYS / "shrid_pc11dist_key.csv", dtype=str)

print(f"names: {len(names):,} rows | cols: {names.columns.tolist()}")
print(f"codes: {len(codes):,} rows | cols: {codes.columns.tolist()}")

# Merge on shrid2, keeping the merge indicator to see what matches.
# NOTE: this is the SHRID-LEVEL merge (~596k rows). It is deliberately named
# shrid_ref to keep it distinct from the collapsed district-level reference
# built later in this script.
shrid_ref = names[["shrid2", "state_name", "district_name"]].merge(
    codes[["shrid2", "pc11_state_id", "pc11_district_id"]],
    on="shrid2", how="outer", indicator=True
)

print("\n--- merge diagnostics ---")
print(shrid_ref["_merge"].value_counts())

# The 4 unmatched: sanity check only
print(shrid_ref[shrid_ref["_merge"] == "right_only"])

# Keep matched rows and collapse to one row per district
pc11_pairs = (shrid_ref[shrid_ref["_merge"] == "both"]
                [["pc11_state_id", "pc11_district_id", "state_name", "district_name"]]
                .drop_duplicates()
                .sort_values(["pc11_state_id", "pc11_district_id"])
                .reset_index(drop=True))

print(f"\nUnique districts: {len(pc11_pairs)}")     # expect ~640

# Integrity check: no district code should carry more than one name
dupes = (pc11_pairs.groupby(["pc11_state_id", "pc11_district_id"])
                   .size()
                   .loc[lambda s: s > 1])
print(f"Codes with multiple names: {len(dupes)}")
if len(dupes):
    print(pc11_pairs.merge(dupes.rename("n").reset_index(),
                           on=["pc11_state_id", "pc11_district_id"]))


# Keep every (code, name) combination as a valid lookup key.
# A single PC11 code may legitimately carry several names because
# districts were split after 2011: SHRUG labels shrids with the current
# district name while retaining the 2011 parent code.
pc11_lookup = (shrid_ref[shrid_ref["_merge"] == "both"]
                 [["pc11_state_id", "pc11_district_id", "state_name", "district_name"]]
                 .drop_duplicates()
                 .sort_values(["pc11_state_id", "pc11_district_id", "district_name"])
                 .reset_index(drop=True))

print(f"Lookup rows (code-name pairs): {len(pc11_lookup)}")          # 696
print(f"Distinct PC11 districts:       "
      f"{pc11_lookup[['pc11_state_id','pc11_district_id']].drop_duplicates().shape[0]}")

# Which codes carry several names, and which names
multi = (pc11_lookup.groupby(["pc11_state_id", "pc11_district_id"])["district_name"]
                    .agg(list)
                    .loc[lambda s: s.str.len() > 1])
print(f"\nCodes with several names: {len(multi)}")
print(multi.head(20))


#####################################################################
# RBI: find all state-district name combinations
#####################################################################

# ---------------------------------------------------------------
# 1. Load the shares file (names and codes are already in memory)
# ---------------------------------------------------------------
shares = pd.read_excel(TABLES / "shares.xlsx",
                       dtype={"pc11_state_id": str, "pc11_district_id": str})

print(f"shares: {len(shares):,} rows | {shares.columns.tolist()}")

# ---------------------------------------------------------------
# 2. Build the code -> name reference
#    Both key files are at shrid level, so district names repeat across
#    thousands of villages. Collapse to ONE name per district code by
#    taking the most frequent name (some codes carry several names
#    because districts were split after 2011).
# ---------------------------------------------------------------
merged = names[["shrid2", "state_name", "district_name"]].merge(
    codes[["shrid2", "pc11_state_id", "pc11_district_id"]],
    on="shrid2", how="inner"
)
print(f"\nAfter shrid-level merge: {len(merged):,} rows")

ref = (merged.groupby(["pc11_state_id", "pc11_district_id",
                       "state_name", "district_name"])
             .size().reset_index(name="n_shrids")
             .sort_values("n_shrids", ascending=False)
             .drop_duplicates(subset=["pc11_state_id", "pc11_district_id"])
             [["pc11_state_id", "pc11_district_id", "state_name", "district_name"]]
             .sort_values(["pc11_state_id", "pc11_district_id"])
             .reset_index(drop=True))

print(f"Districts in reference: {len(ref)}")        # expect 640


#####################################################################
# 2b. Resolve district names shared by more than one PC11 code
#####################################################################

# `ref` holds one row per PC11 code, but several codes share the SAME name.
# The clearest case is Mumbai: SHRUG labels both 518 (Mumbai Suburban) and
# 519 (Mumbai City) as "mumbai suburban". Any later match done on names would
# return two rows for one district and silently duplicate it downstream.
# List every affected name first, so each case is resolved deliberately.
dup_all = (ref.groupby(["state_name", "district_name"])
              .size()
              .loc[lambda s: s > 1])
print(f"\nNames shared by more than one PC11 code: {len(dup_all)}")
if len(dup_all):
    print(ref.merge(dup_all.rename("n_codes").reset_index(),
                    on=["state_name", "district_name"])
             .sort_values(["state_name", "district_name"])
             .to_string())

# Mumbai: in the raw RBI directory code 519 holds 2,776 branches while 518
# holds exactly one, so 519 is where the banking data actually sits. Drop 518
# from the lookup so the name resolves to a single code.
ref_lookup = ref[~((ref["pc11_state_id"] == "27") &
                   (ref["pc11_district_id"] == "518"))].copy()


# Delhi cannot be used: 95.3% of its branches carry no district code at all,
# and SHRUG labels all nine remaining district codes (090-098) with the same
# name, "north west". Neither problem is fixable from this source, so Delhi
# is dropped from the lookup to prevent a name match from returning nine rows.
ref_lookup = ref_lookup[ref_lookup["pc11_state_id"] != "07"].copy()

dup_names = (ref_lookup.groupby(["state_name", "district_name"])
                       .size().loc[lambda s: s > 1])
print(f"Remaining duplicated names: {len(dup_names)}")   # want 0
if len(dup_names):
    print(dup_names)


# ---------------------------------------------------------------
# 3. Attach names to the shares file
# ---------------------------------------------------------------
# Harmonise ID formats first (Excel can strip leading zeros)
shares["pc11_state_id"]    = shares["pc11_state_id"].str.strip().str.zfill(2)
shares["pc11_district_id"] = shares["pc11_district_id"].str.strip().str.zfill(3)
ref["pc11_state_id"]       = ref["pc11_state_id"].str.strip().str.zfill(2)
ref["pc11_district_id"]    = ref["pc11_district_id"].str.strip().str.zfill(3)
# ref_lookup is the de-duplicated copy used later for the crosswalk, so it
# needs the same ID formatting.
ref_lookup["pc11_state_id"]    = ref_lookup["pc11_state_id"].str.strip().str.zfill(2)
ref_lookup["pc11_district_id"] = ref_lookup["pc11_district_id"].str.strip().str.zfill(3)

shares_named = shares.merge(ref, on=["pc11_state_id", "pc11_district_id"], how="left", indicator=True)

print(f"\nRows before: {len(shares):,} | after: {len(shares_named):,}")   # must be equal
print(shares_named["_merge"].value_counts())
print(f"Rows without a name: {shares_named['district_name'].isna().sum()}")

shares_named = shares_named.drop(columns="_merge")
shares_named.to_excel(TABLES / "shares with names.xlsx", index=False)
print(f"\nSaved: {TABLES / 'shares with names.xlsx'}")

print(f"Rows: {len(shares_named):,}")
print(f"Rows without a name: {shares_named['district_name'].isna().sum()}")   # want 0
print(f"Distinct districts with names: "
      f"{shares_named[['state_name','district_name']].drop_duplicates().shape[0]}")
print(shares_named.head(10))


# ---------------------------------------------------------------
# 4. Unique state-district name pairs on the shares side
# ---------------------------------------------------------------
shares_pairs = (shares_named[["state_name", "district_name"]]
                  .dropna()
                  .drop_duplicates()
                  .sort_values(["state_name", "district_name"])
                  .reset_index(drop=True))

print(f"Unique state-district pairs (shares): {len(shares_pairs)}")   # expect 632
print(shares_pairs.head(20))

out_path = TABLES / "shares state-district pairs.xlsx"
shares_pairs.to_excel(out_path, index=False)
print(f"Saved to: {out_path}")

#####################################################################
# MATCHING: load the two pair lists
#####################################################################

cphs_pairs   = pd.read_excel(TABLES / "CPHS state-district pairs.xlsx", dtype=str)
shares_pairs = pd.read_excel(TABLES / "shares state-district pairs.xlsx", dtype=str)

print(f"CPHS pairs:   {len(cphs_pairs):>4} | cols: {cphs_pairs.columns.tolist()}")
print(f"Shares pairs: {len(shares_pairs):>4} | cols: {shares_pairs.columns.tolist()}")


#####################################################################
# MATCHING: load the Local Government Directory (second reference)
#####################################################################

# LGD is the official directory of Indian administrative units, maintained by
# the Ministry of Panchayati Raj with the Registrar General of India. It gives
# the Census 2011 code of each CURRENT district, archives superseded names, and
# requires a government order for every change to the directory.
#
# It is the natural complement to SHRUG here: SHRUG carries the 2011 Census
# names (the OLD ones), LGD carries today's names (the NEW ones), and CPHS uses
# a mixture of both. Matching against each in turn resolves most naming
# differences from published sources rather than by hand.
#
# Downloaded from lgdirectory.gov.in > "LGD Codes of Districts" > All States.
# Keep the file with the project: LGD is a live directory, so the download date
# is the version identifier.
# NOTE: the portal exports .xlsx with malformed styles that openpyxl rejects,
# hence engine="calamine" (pip install python-calamine).
LGD_FILE = DATA / "LGD - Local Government Directory, Government of India.xlsx"

lgd = pd.read_excel(LGD_FILE, dtype=str, engine="calamine")
print(f"\nLGD districts: {len(lgd):,}")

# Districts created after 2011 carry no census code (recorded as "000") and
# cannot serve as evidence, so they are dropped.
lgd = lgd[lgd["Census2011 Code"] != "000"].copy()
lgd["pc11_district_id"] = lgd["Census2011 Code"].str.strip().str.zfill(3)
print(f"LGD districts with a Census 2011 code: {len(lgd):,}")


#####################################################################
# MATCHING: normalise names on all three sources
#####################################################################

def norm(s):
    """Uppercase, strip punctuation and extra spacing, and drop the connective
    'AND'. The last step matters: CPHS writes 'Jammu & Kashmir' while LGD writes
    'Jammu And Kashmir', and without it every district in that state fails to
    match on the state key alone."""
    if pd.isna(s):
        return ""
    s = str(s).upper().strip()
    for ch in [".", ",", "'", "`", "-", "&", "(", ")", "/"]:
        s = s.replace(ch, " ")
    s = " ".join(s.split())
    return " ".join(t for t in s.split() if t != "AND")


# --- CPHS side ---
cphs_m = (cphs_pairs[["state", "district"]]
            .rename(columns={"state": "state_cphs", "district": "district_cphs"})
            .copy())

# Telangana was created in June 2014 and does not exist in the 2011 Census, so
# its districts sit under Andhra Pradesh in PC11. Delhi is a naming convention
# only. Applied to the ORIGINAL strings, before normalisation.
STATE_RECODE = {
    "Telangana": "Andhra Pradesh",
    "Delhi":     "NCT of Delhi",
}
cphs_m["state_fix"] = cphs_m["state_cphs"].replace(STATE_RECODE)
cphs_m["state_n"]   = cphs_m["state_fix"].map(norm)
cphs_m["dist_n"]    = cphs_m["district_cphs"].map(norm)

n_state_recoded = (cphs_m["state_cphs"] != cphs_m["state_fix"]).sum()
print(f"\nDistricts recoded to a pre-2011 state: {n_state_recoded}")

# --- SHRUG side (PC11 names + codes), de-duplicated ---
ref_lookup["state_n"] = ref_lookup["state_name"].map(norm)
ref_lookup["dist_n"]  = ref_lookup["district_name"].map(norm)

# --- LGD side (current names + census codes) ---
# LGD state names carry a "(State)" / "(UT)" suffix that must be stripped.
# Telangana is a separate state in LGD but its districts keep Andhra Pradesh's
# 2011 census codes, so the same STATE_RECODE is applied here.
lgd["state_clean"] = (lgd["State Name"]
                        .str.replace("(State)", "", regex=False)
                        .str.replace("(UT)", "", regex=False)
                        .str.strip())
lgd["state_n"] = lgd["state_clean"].replace(STATE_RECODE).map(norm)
lgd["dist_n"]  = lgd["District Name (In English)"].map(norm)

lgd_lookup = (lgd[["state_n", "dist_n", "pc11_district_id",
                   "District Name (In English)"]]
                .rename(columns={"District Name (In English)": "lgd_name"})
                .drop_duplicates())

# Delhi is excluded from BOTH references, not just SHRUG. LGD does carry valid
# 2011 codes for Delhi's districts, but the RBI directory assigns no district
# to 95.3% of Delhi's branches, so a share computed for those codes would rest
# on almost no data. Letting the LGD layer assign a code here would produce a
# treatment value that looks valid and is not.
lgd_lookup = lgd_lookup[lgd_lookup["state_n"] != norm("NCT of Delhi")].copy()

#####################################################################
# LAYER 1: exact match against SHRUG (2011 Census names)
#####################################################################

# Catches every district whose CPHS name still uses the 2011 spelling. This is
# purely mechanical: normalisation only, no judgement.
l1 = cphs_m.merge(
    ref_lookup[["state_n", "dist_n", "pc11_state_id", "pc11_district_id"]],
    on=["state_n", "dist_n"], how="left")
l1["source"] = l1["pc11_district_id"].notna().map({True: "SHRUG (2011 name)", False: ""})

print(f"\nLAYER 1 - exact vs SHRUG: {l1['pc11_district_id'].notna().sum()} of {len(l1)}")


#####################################################################
# LAYER 2: exact match against LGD (current names)
#####################################################################

# Catches districts that CPHS reports under their POST-2011 name, which SHRUG
# cannot know: Prayagraj (Allahabad), Gurugram (Gurgaon), Bengaluru Urban
# (Bangalore), Kalaburagi (Gulbarga), and so on. Resolved from a government
# directory rather than by hand.
todo = l1[l1["pc11_district_id"].isna()].drop(
    columns=["pc11_state_id", "pc11_district_id", "source"])

l2 = todo.merge(lgd_lookup[["state_n", "dist_n", "pc11_district_id"]],
                on=["state_n", "dist_n"], how="left")
l2 = l2.drop_duplicates(subset=["state_cphs", "district_cphs"])
l2["source"] = l2["pc11_district_id"].notna().map({True: "LGD (current name)", False: ""})

# LGD has no state code column, so recover it from the SHRUG reference.
state_codes = ref_lookup[["state_n", "pc11_state_id"]].drop_duplicates()
l2 = l2.merge(state_codes, on="state_n", how="left")

print(f"LAYER 2 - exact vs LGD:   {l2['pc11_district_id'].notna().sum()} of {len(l2)}")


#####################################################################
# LAYER 3: fuzzy candidates against LGD, for review
#####################################################################

# What remains are spelling differences that neither reference spells the way
# CPHS does (Khurda/Khordha, Keonjhar/Kendujhar, Angul/Anugola). Fuzzy matching
# is used to PROPOSE a candidate with a similarity score, not to decide: at low
# scores it is unreliable, and even at high scores it can be wrong (it scores
# "North Bastar (Kanker)" against plain "Bastar" at 90 while the correct
# district is "Uttar Bastar Kanker"). Every candidate here is confirmed by hand
# in DISTRICT_RECODE below, and the score is exported so a reviewer can see the
# evidence each decision rests on.
#
# Candidates are restricted to the same state, which is essential: several
# district names repeat across states (Aurangabad, Bilaspur, Bijapur).
from rapidfuzz import process, fuzz

resid = l2[l2["pc11_district_id"].isna()].copy()

cand_rows = []
for _, r in resid.iterrows():
    pool = lgd_lookup[lgd_lookup["state_n"] == r["state_n"]]
    if pool.empty:
        cand_rows.append({"state_cphs": r["state_cphs"], "district_cphs": r["district_cphs"],
                          "lgd_candidate": None, "score": 0.0, "runner_up_score": 0.0,
                          "candidate_code": None})
        continue
    hits = process.extract(r["dist_n"], pool["dist_n"].tolist(),
                           scorer=fuzz.WRatio, limit=2)
    best, score, idx = hits[0]
    runner = hits[1][1] if len(hits) > 1 else 0.0
    cand_rows.append({"state_cphs": r["state_cphs"], "district_cphs": r["district_cphs"],
                      "lgd_candidate": pool.iloc[idx]["lgd_name"],
                      "score": round(score, 1), "runner_up_score": round(runner, 1),
                      "candidate_code": pool.iloc[idx]["pc11_district_id"]})

candidates = pd.DataFrame(cand_rows).sort_values("score", ascending=False)
print(f"\nLAYER 3 - fuzzy candidates to confirm: {len(candidates)}")
print(candidates.to_string(index=False))

# NOTE: not exported separately. The candidate, its score and whether it
# agreed with the final decision are merged into the crosswalk below, so the
# reviewer works from a single file.


#####################################################################
# LAYER 4: confirmed manual resolutions
#####################################################################

# Only the residual from Layer 3 appears here. Each entry was checked against
# the LGD candidate above and against the district's entry in LGD, which links
# the government order behind each renaming. Four of these DISAGREE with the
# fuzzy candidate, which is why the fuzzy step cannot be trusted on its own:
#   North Bastar (Kanker) - fuzzy proposes "Bastar" (414); correct is
#                           "Uttar Bastar Kanker" (413)
#   Sonepur               - fuzzy proposes "Puri" (387); correct is
#                           "Subarnapur" (392)
#   Angul                 - fuzzy proposes "Nabarangpur" (397); correct is
#                           "Anugola" (384)
#   West Sikkim           - fuzzy proposes "Gangtok" (244); correct is
#                           "Gyalshing" (242)
# Sikkim renamed its directional districts in 2021, so none of the three has
# any textual similarity to its current name.
DISTRICT_RECODE = {
    # (normalised state, normalised CPHS district): PC11 district code
    ("ANDHRA PRADESH", "Y S R"):                 "551",   # LGD: Y.S.R. Kadapa
    ("CHHATTISGARH",   "DANTEWADA"):             "416",   # LGD: Dakshin Bastar Dantewada
    ("CHHATTISGARH",   "KABIRDHAM"):             "407",   # LGD: Kabeerdham
    ("CHHATTISGARH",   "NORTH BASTAR KANKER"):   "413",   # LGD: Uttar Bastar Kanker
    ("JHARKHAND",      "DEOGARH"):               "350",   # LGD: Deoghar
    ("KARNATAKA",      "YADGIRI"):               "580",   # LGD: Yadgir
    ("MADHYA PRADESH", "EAST NIMAR KHANDWA"):    "466",   # LGD: Khandwa (East Nimar)
    ("MADHYA PRADESH", "WEST NIMAR KHARGONE"):   "440",   # LGD: Khargone (West Nimar)
    ("ODISHA",         "ANGUL"):                 "384",   # LGD: Anugola
    ("ODISHA",         "DEOGARH"):               "373",   # LGD: Debagada
    ("ODISHA",         "JAGATSINGHPUR"):         "380",   # LGD: Jagatsinghapur
    ("ODISHA",         "KEONJHAR"):              "375",   # LGD: Kendujhar
    ("ODISHA",         "KHURDA"):                "386",   # LGD: Khordha
    ("ODISHA",         "SONEPUR"):               "392",   # LGD: Subarnapur
    ("SIKKIM",         "EAST SIKKIM"):           "244",   # LGD: Gangtok (renamed 2021)
    ("SIKKIM",         "SOUTH SIKKIM"):          "243",   # LGD: Namchi (renamed 2021)
    ("SIKKIM",         "WEST SIKKIM"):           "242",   # LGD: Gyalshing (renamed 2021)
    ("UTTAR PRADESH",  "MAHARAJGANJ"):           "187",   # LGD: Mahrajganj
    ("UTTAR PRADESH",  "SHRAVASTI"):             "181",   # LGD: Shrawasti
}

l3 = resid.drop(columns=["pc11_district_id", "source"]).copy()
l3["pc11_district_id"] = [DISTRICT_RECODE.get((s, d))
                          for s, d in zip(l3["state_n"], l3["dist_n"])]
l3["source"] = l3["pc11_district_id"].notna().map(
    {True: "Manual (confirmed vs LGD)", False: "No match"})
# NOTE: no state_codes merge here. `resid` is derived from l2, which already
# carries pc11_state_id; merging again would produce _x/_y suffixed columns
# and break the concat below.

print(f"\nLAYER 4 - manual, confirmed: {l3['pc11_district_id'].notna().sum()} of {len(l3)}")
print(f"Still unmatched: {l3['pc11_district_id'].isna().sum()}")


#####################################################################
# ASSEMBLE the crosswalk
#####################################################################

cols = ["state_cphs", "district_cphs", "state_n", "dist_n",
        "pc11_state_id", "pc11_district_id", "source"]

crosswalk = pd.concat([
    l1[l1["pc11_district_id"].notna()][cols],
    l2[l2["pc11_district_id"].notna()][cols],
    l3[cols],
], ignore_index=True)

# The merge must not create rows: one CPHS pair in, one row out.
print(f"\nInput pairs: {len(cphs_m)} | crosswalk rows: {len(crosswalk)}")   # must be equal
print(crosswalk["source"].value_counts())

# Attach the names from both references as evidence for the reviewer
crosswalk = crosswalk.merge(
    ref_lookup[["pc11_state_id", "pc11_district_id", "state_name", "district_name"]]
        .rename(columns={"state_name": "PC11 state name",
                         "district_name": "PC11 district name"}),
    on=["pc11_state_id", "pc11_district_id"], how="left")

lgd_by_code = (lgd.groupby("pc11_district_id")["District Name (In English)"]
                  .agg(lambda s: " / ".join(sorted(set(s))))
                  .rename("LGD district name (current)"))
crosswalk = crosswalk.merge(lgd_by_code, left_on="pc11_district_id",
                            right_index=True, how="left")

crosswalk_out = (crosswalk[["state_cphs", "district_cphs",
                            "PC11 state name", "PC11 district name",
                            "LGD district name (current)",
                            "pc11_state_id", "pc11_district_id", "source"]]
                   .rename(columns={"state_cphs":       "CPHS state",
                                    "district_cphs":    "CPHS district",
                                    "pc11_state_id":    "PC11 state code",
                                    "pc11_district_id": "PC11 district code",
                                    "source":           "Resolved by"})
                   .sort_values(["CPHS state", "CPHS district"])
                   .reset_index(drop=True))

# Attach the fuzzy score to the manually confirmed rows, so the reviewer sees
# how close the automatic candidate was and where it disagreed.
crosswalk_out = crosswalk_out.merge(
    candidates.rename(columns={"state_cphs": "CPHS state",
                               "district_cphs": "CPHS district",
                               "lgd_candidate": "Fuzzy candidate (LGD)",
                               "score": "Fuzzy score",
                               "candidate_code": "Fuzzy candidate code"})
              [["CPHS state", "CPHS district", "Fuzzy candidate (LGD)",
                "Fuzzy score", "Fuzzy candidate code"]],
    on=["CPHS state", "CPHS district"], how="left")

crosswalk_out["Fuzzy agrees"] = ""
mask = crosswalk_out["Fuzzy candidate code"].notna()
crosswalk_out.loc[mask, "Fuzzy agrees"] = (
    crosswalk_out.loc[mask, "Fuzzy candidate code"] ==
    crosswalk_out.loc[mask, "PC11 district code"]).map({True: "Yes", False: "NO"})

# A blank code is a data limitation, not an oversight.
crosswalk_out["Note"] = ""
crosswalk_out.loc[crosswalk_out["Resolved by"] == "No match", "Note"] = (
    "Delhi: no usable district code in the RBI directory via SHRUG")

crosswalk_out["Correct? (Y/N)"]   = ""
crosswalk_out["Reviewer comment"] = ""

# Final column order: the two source names first so they can be compared at a
# glance, then the codes, then how the match was made and the fuzzy evidence,
# then the reviewer's own columns.
crosswalk_out = crosswalk_out[[
    "CPHS state", "CPHS district",
    "PC11 state name", "PC11 district name", "LGD district name (current)",
    "PC11 state code", "PC11 district code",
    "Resolved by",
    "Fuzzy candidate (LGD)", "Fuzzy score", "Fuzzy agrees",
    "Note", "Correct? (Y/N)", "Reviewer comment"]]

# ---------------------------------------------------------------
# Export: one workbook, two sheets (Notes + Crosswalk)
# ---------------------------------------------------------------
NOTES = pd.DataFrame({
 "Item": [
   "Purpose", "Source A - CPHS", "Source B - RBI branch data",
   "Reference 1 - SHRUG (2011 Census names)",
   "Reference 2 - LGD (current names)",
   "Method: four layers",
   'Column "Resolved by"',
   'Columns "Fuzzy candidate / score / agrees"',
   "Result",
   "Districts not matched",
   "Why the state matters",
   "Several CPHS names, one PC11 code",
   "Mumbai",
   "What we would like reviewed",
 ],
 "Detail": [
   "Map CPHS district names onto 2011 Census (PC11) district codes, so the RBI branch shares can be merged onto the CPHS household panel.",
   "574 unique state-district pairs, extracted from all 144 monthly CPHS parquet files (Jan 2014 - Dec 2025). The full panel was used rather than a single month because geographic coverage grows over the period, from about 394 districts in early 2014 to about 526 by the end.",
   "RBI Bank Branch Directory via SHRUG v2.1. It carries PC11 district codes but no district names.",
   "SHRUG core keys (shrid_loc_names + shrid_pc11dist_key), collapsed to one row per district. These are the district names as they stood in the 2011 Census.",
   "Local Government Directory, Ministry of Panchayati Raj, lgdirectory.gov.in ('LGD Codes of Districts' > All States). LGD gives the Census 2011 code of each CURRENT district, archives superseded names, and requires a government order for every change to the directory. It is the complement to SHRUG: SHRUG has the old names, LGD the new ones, and CPHS uses a mixture of both.",
   "1) Exact match against SHRUG, after normalising case, punctuation and spacing. 2) Exact match against LGD, catching districts CPHS reports under their post-2011 name. 3) Fuzzy string matching against LGD to PROPOSE a candidate for whatever remained. 4) Manual confirmation of each candidate against its LGD entry.",
   "Which layer resolved the row. 'SHRUG (2011 name)' and 'LGD (current name)' are mechanical exact matches with no judgement involved. 'Manual (confirmed vs LGD)' means the row was resolved by hand and is where review is most useful.",
   "For the manually resolved rows only: the district the fuzzy step proposed, its similarity score, and whether it agreed with the final decision. 'Fuzzy agrees' = NO marks the cases where the automatic candidate was wrong and was overridden - which is why the fuzzy step is used to propose rather than to decide.",
   "561 of 574 CPHS pairs received a PC11 code (97.7%). The crosswalk is one-to-one: 574 pairs in, 574 rows out, so no household is duplicated when it is applied to the panel.",
   "All are Delhi. This is a data limitation, not a naming problem: 95.3% of Delhi's branches carry no district code at all in the RBI directory ('000'), and SHRUG labels the nine remaining Delhi district codes (090-098) identically. Whether to exclude Delhi from the analysis, or to recover district assignment from the RBI's primary publication, is a decision we would like to discuss.",
   "All matching is keyed on state AND district, because the same name can refer to different districts in different states. 'Deogarh' appears twice: it is Deoghar in Jharkhand and Debagada in Odisha. Both are correct.",
   "Some PC11 codes are reached by two CPHS names. This is expected: CPHS uses two spellings over time (Anantapur / Ananthapuramu), and districts split after 2011 map back to their 2011 parent. All were inspected and none is a collision between genuinely different districts.",
   "CPHS reports a single 'Mumbai'. PC11 splits the city into Mumbai City (519) and Mumbai Suburban (518), but SHRUG labels both codes 'mumbai suburban', and the RBI directory places 2,776 branches under 519 against a single branch under 518. CPHS 'Mumbai' was mapped to 519.",
   "Please confirm the mappings using the 'Correct? (Y/N)' and 'Reviewer comment' columns. Filtering 'Resolved by' = 'Manual (confirmed vs LGD)' groups the rows that most need attention; filtering 'Fuzzy agrees' = NO isolates the few where the automatic candidate had to be overridden.",
 ]})

out_path = TABLES / "CPHS to PC11 crosswalk.xlsx"
with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
    NOTES.to_excel(writer, sheet_name="Notes", index=False)
    crosswalk_out.to_excel(writer, sheet_name="Crosswalk", index=False)

    # Readable column widths, wrapped notes, frozen header and autofilter
    from openpyxl.styles import Alignment
    from openpyxl.utils import get_column_letter

    ws = writer.sheets["Notes"]
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 118
    for r in range(1, ws.max_row + 1):
        for c in (1, 2):
            ws.cell(row=r, column=c).alignment = Alignment(vertical="top",
                                                           wrap_text=True)

    ws2 = writer.sheets["Crosswalk"]
    for i, col in enumerate(crosswalk_out.columns, start=1):
        lens = crosswalk_out[col].dropna().astype(str).map(len)
        width = max(len(str(col)), int(lens.max()) if len(lens) else 0)
        ws2.column_dimensions[get_column_letter(i)].width = min(width + 3, 50)
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

print(f"\nSaved: {out_path}")
print("\nResolved by:")
print(crosswalk_out["Resolved by"].value_counts())


#####################################################################
# FINAL CHECKS
#####################################################################

# Several CPHS names can legitimately share one PC11 code: two spellings used
# over time (Anantapur / Ananthapuramu), or a post-2011 split mapping back to
# its 2011 parent. What must NOT happen is a collision between two genuinely
# different districts, so print them all to be eyeballed.
multi = (crosswalk_out[crosswalk_out["Resolved by"] != "No match"]
           .groupby(["PC11 state code", "PC11 district code"])
           .agg(n_names=("CPHS district", "nunique"),
                names=("CPHS district", lambda s: sorted(set(s))))
           .reset_index()
           .loc[lambda d: d["n_names"] > 1])

print(f"\nPC11 codes reached by more than one CPHS name: {len(multi)}")
print(multi.to_string(index=False))

# Line ~656: check whether Delhi slipped through the LGD layer
print(crosswalk_out[crosswalk_out["CPHS state"] == "Delhi"]
        [["CPHS district", "PC11 district code", "Resolved by"]].to_string())